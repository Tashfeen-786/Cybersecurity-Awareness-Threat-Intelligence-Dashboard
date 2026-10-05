"""Render reports/project_report.md into reports/project_report.docx.

Run from the project root:  python scripts/make_docx_report.py
(Requires python-docx; the markdown file is the source of truth.)
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "reports" / "project_report.md"
OUT = ROOT / "reports" / "project_report.docx"

DARK = RGBColor(0x1B, 0x2A, 0x41)
ACCENT = RGBColor(0x0E, 0x6E, 0x9E)
GREY = RGBColor(0x55, 0x5F, 0x6B)


def add_runs(par, text: str) -> None:
    """Render **bold** / *italic* / `code` inline markdown."""
    for chunk in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            par.add_run(chunk[1:-1]).italic = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = par.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xB3, 0x3A, 0x3A)
        else:
            par.add_run(chunk)


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for section in doc.sections:
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    lines = MD.read_text(encoding="utf-8").split("\n")
    i = 0
    in_code = False
    code_buf: list[str] = []

    while i < len(lines):
        ln = lines[i]

        if ln.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                p.paragraph_format.left_indent = Inches(0.25)
                code_buf = []
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code_buf.append(ln)
            i += 1
            continue

        if not ln.strip():
            i += 1
            continue

        if ln.startswith("# "):          # title
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(ln[2:].strip())
            r.bold = True
            r.font.size = Pt(19)
            r.font.color.rgb = DARK
        elif ln.startswith("## "):       # section
            p = doc.add_heading(level=1)
            r = p.add_run(ln[3:].strip())
            r.font.size = Pt(13)
            r.font.color.rgb = ACCENT
            r.bold = True
        elif ln.startswith("### "):
            p = doc.add_heading(level=2)
            r = p.add_run(ln[4:].strip())
            r.font.size = Pt(11.5)
            r.font.color.rgb = DARK
            r.bold = True
        elif ln.strip() == "---":
            doc.add_paragraph()
        elif ln.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, ln[2:].strip())
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = GREY
        elif ln.lstrip().startswith("|") and i + 1 < len(lines) and \
                set(lines[i + 1].replace("|", "").replace(" ", "")) <= {"-", ":"}:
            # markdown table
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not set("".join(cells)) <= {"-", ":", " "}:
                    rows.append(cells)
                i += 1
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row[:len(rows[0])]):
                        c = table.cell(ri, ci)
                        c.text = ""
                        para = c.paragraphs[0]
                        add_runs(para, cell)
                        for r in para.runs:
                            r.font.size = Pt(9)
                            if ri == 0:
                                r.bold = True
                doc.add_paragraph()
            continue
        elif re.match(r"^\s*[-*] ", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, ln.split(" ", 1)[1])
        elif re.match(r"^\s*\d+\. ", ln):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.split(r"^\s*\d+\. ", ln)[1])
        else:
            p = doc.add_paragraph()
            add_runs(p, ln.strip())
        i += 1

    doc.save(OUT)
    print(f"written: {OUT}")


if __name__ == "__main__":
    main()
